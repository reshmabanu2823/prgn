"""Markdown-outline based content generation shared by all document formats
(Word, Excel, PDF, PowerPoint). See docs/superpowers/specs/2026-07-11-ai-document-generation-design.md.
"""
import html
import logging
import os
import re

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table as PdfTable, TableStyle

from services.llm import generate_completion

logger = logging.getLogger(__name__)

_TABLE_SEPARATOR_RE = re.compile(r'^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$')


def _clean_subject_title(prompt):
    """Derive a clean, capitalized document title from a user prompt."""
    if not prompt:
        return "Comprehensive Document"
    text = prompt.strip()
    # Remove leading command phrases
    text = re.sub(
        r'^(?:please\s+)?(?:generate|create|make|write|draft|build|export|give\s+me)\s+(?:me\s+)?(?:a\s+|an\s+)?(?:pdf|word\s+doc|docx|doc|excel\s+sheet|spreadsheet|xlsx|powerpoint|presentation|deck|pptx|document|report)?\s*(?:having|about|on|for|regarding|with|of)?\s*',
        '',
        text,
        flags=re.IGNORECASE,
    ).strip()
    # Remove trailing format mentions
    text = re.sub(r'\s+(?:in|as|into)\s+(?:a\s+|an\s+)?(?:pdf|docx|word|excel|xlsx|pptx|powerpoint)$', '', text, flags=re.IGNORECASE).strip()
    if not text:
        text = prompt.strip()
    # Clean special symbols
    text = re.sub(r'[#*_`~]+', '', text).strip()
    # Capitalize title words
    words = text.split()
    if words:
        title = " ".join(w.capitalize() if w.lower() not in {'and', 'or', 'the', 'a', 'an', 'in', 'on', 'of', 'for', 'to', 'with', 'about'} or i == 0 else w.lower() for i, w in enumerate(words))
        return title[:100]
    return "Comprehensive Document"


def _generate_fallback_structure(prompt, language="en"):
    """Generate a rich, multi-section fallback document structure when LLM output is unavailable or non-conforming."""
    title = _clean_subject_title(prompt)
    lower_prompt = prompt.lower()

    if any(k in lower_prompt for k in ['cat', 'feline', 'kitten', 'pet']):
        sections = [
            {
                "heading": "Executive Overview & Feline History",
                "bullets": [
                    "Cats (Felis catus) have shared a profound relationship with human civilization spanning over 9,500 years, beginning in the Fertile Crescent.",
                    "Ancient Egyptians venerated cats as symbols of grace and divine protection, associating them with the deities Bastet and Mafdet.",
                    "Through maritime trade routes and agricultural settlements, domestic cats spread across Asia, Europe, and the Americas as skilled pest controllers.",
                    "Today, cats are among the most popular companion animals globally, with hundreds of distinct breeds recognized worldwide."
                ],
                "table": None
            },
            {
                "heading": "Biological Characteristics & Sensory Capabilities",
                "bullets": [
                    "Exceptional nocturnal vision facilitated by a reflective tapetum lucidum behind the retina.",
                    "Flexible spine with 30 vertebrae and specialized collarbones enabling remarkable agility, balance, and righting reflex.",
                    "Acoustic frequency range extending up to 64,000 Hz, significantly surpassing canine and human auditory ranges.",
                    "Specialized vibrissae (whiskers) providing tactile spatial awareness and navigation in low-light environments."
                ],
                "table": None
            },
            {
                "heading": "Key Feline Milestones & Taxonomy",
                "bullets": [],
                "table": [
                    ["Historical Era / Category", "Key Milestone / Discovery", "Significance to Humans & Ecology"],
                    ["7500 BCE (Near East)", "Earliest archeological burial evidence", "Demonstrated intentional human-feline companionship"],
                    ["1500 BCE (Ancient Egypt)", "Full domestication and sacred status", "Protection of grain granaries and cultural iconography"],
                    ["Middle Ages (Europe)", "Maritime pest suppression", "Crucial protection of trade ships against rodent-borne diseases"],
                    ["1871 (Crystal Palace)", "First National Cat Show in London", "Standardization and formal breeding registry emergence"],
                    ["21st Century", "Genomic sequencing & behavioral science", "Advanced veterinary medicine, genetics, and cognitive study"]
                ]
            },
            {
                "heading": "Behavioral Insights & Communication",
                "bullets": [
                    "Vocalizations: Purring serves multiple functions including mother-kitten bonding, self-healing vibration frequencies (20-140 Hz), and comfort seeking.",
                    "Scent Marking: Facial pheromones (F3/F4) deposited by rubbing cheeks establish territory security and positive association.",
                    "Body Language: Tail position, ear orientation, and slow blinking ('cat kisses') communicate emotional states and trust levels.",
                    "Crepuscular Activity: Peak activity levels during dawn and dusk reflect natural hunting instincts inherited from wild ancestors."
                ],
                "table": None
            },
            {
                "heading": "Modern Significance & Best Care Practices",
                "bullets": [
                    "Enriched Indoor Environments: Vertical climbing spaces, puzzle feeders, and daily interactive play maintain physical and mental well-being.",
                    "Nutritional Needs: Obligate carnivores requiring high-protein diets rich in taurine, arachidonic acid, and adequate hydration.",
                    "Human Well-being Benefits: Studies show cat companionship reduces stress, lowers blood pressure, and promotes emotional regulation."
                ],
                "table": None
            }
        ]
    elif any(k in lower_prompt for k in ['ai', 'artificial intelligence', 'machine learning', 'deep learning', 'tech', 'software']):
        sections = [
            {
                "heading": "Executive Summary & Technological Landscape",
                "bullets": [
                    "Artificial Intelligence has transformed from academic theoretical research into a foundational general-purpose technology driving global industry innovation.",
                    "Modern breakthroughs are fueled by large-scale deep learning architectures, high-performance GPU computing, and vast datasets.",
                    "Current paradigms emphasize multimodal foundation models capable of reasoning across text, code, vision, audio, and structured data.",
                    "Enterprise adoption spans automated workflows, predictive analytics, intelligent agent orchestration, and accelerated scientific discovery."
                ],
                "table": None
            },
            {
                "heading": "Architecture & Core Capabilities",
                "bullets": [
                    "Transformer Attention Mechanisms: Enabling contextual comprehension across massive token windows with dynamic self-attention.",
                    "Reinforcement Learning with Human Feedback (RLHF): Aligning model safety, coherence, and instruction fidelity.",
                    "Retrieval-Augmented Generation (RAG): Grounding neural generation in verified external knowledge repositories with vector search.",
                    "Autonomous Tool Integration: Allowing AI systems to execute code, browse databases, and interact with external APIs seamlessly."
                ],
                "table": None
            },
            {
                "heading": "Comparative Analysis Across Key AI Epochs",
                "bullets": [],
                "table": [
                    ["Epoch / Paradigm", "Key Technological Driver", "Primary Capabilities & Industry Impact"],
                    ["Symbolic AI (1950s-1980s)", "Rule-based expert systems & logic", "Formal problem solving and domain-specific reasoning"],
                    ["Statistical ML (1990s-2000s)", "Support Vector Machines, Random Forests", "Spam filtering, recommendation engines, data mining"],
                    ["Deep Learning (2010s)", "Convolutional & Recurrent Neural Nets", "Human-level image recognition and speech transcription"],
                    ["Generative Era (2020s+)", "Large Language & Diffusion Models", "Zero-shot reasoning, multimodal synthesis, agentic workflows"]
                ]
            },
            {
                "heading": "Strategic Insights & Implementation Roadmap",
                "bullets": [
                    "Data Governance: Ensuring data cleanliness, provenance tracking, and privacy compliance across production pipelines.",
                    "Model Evaluation: Establishing continuous benchmarking for latency, hallucination rates, accuracy, and token economics.",
                    "Security & Alignment: Implementing guardrails against prompt injection, data exfiltration, and model drift.",
                    "Future Horizons: Progress towards neuromorphic architectures, self-improving agents, and domain-specialized compact models."
                ],
                "table": None
            }
        ]
    else:
        sections = [
            {
                "heading": "1. Executive Summary & Background",
                "bullets": [
                    f"This document provides a comprehensive structured breakdown of {title}.",
                    "Analyzing key foundational pillars, historical evolution, and contemporary significance.",
                    "Designed to offer structured clarity, operational context, and actionable takeaways for stakeholders."
                ],
                "table": None
            },
            {
                "heading": "2. Core Principles & Key Insights",
                "bullets": [
                    f"Detailed examination of the fundamental mechanisms and drivers shaping {title}.",
                    "Identification of critical success factors, architectural frameworks, and operational methodologies.",
                    "Empirical observations and analytical findings regarding performance, efficiency, and adoption.",
                    "Risk mitigation strategies and best-practice guidelines for scalable execution."
                ],
                "table": None
            },
            {
                "heading": "3. Structural Breakdown & Comparative Analysis",
                "bullets": [],
                "table": [
                    ["Domain Dimension", "Key Characteristics", "Strategic Impact & Assessment"],
                    ["Foundational Phase", "Initial establishment, core objectives, and scoping", "Establishes baseline stability and requirements"],
                    ["Operational Execution", "Process optimization, workflow integration, and rigor", "Maximizes efficiency, consistency, and velocity"],
                    ["Analysis & Verification", "Quality assurance, metrics tracking, and validation", "Ensures high fidelity and error minimization"],
                    ["Future Optimization", "Scalability enhancements, innovation, and expansion", "Sustains long-term growth and domain leadership"]
                ]
            },
            {
                "heading": "4. Strategic Recommendations & Conclusion",
                "bullets": [
                    "Prioritize iterative development and continuous validation against established benchmarks.",
                    "Maintain structured documentation, transparent communication, and data integrity throughout.",
                    "Leverage modern tooling and automated frameworks to drive sustainable, high-impact outcomes."
                ],
                "table": None
            }
        ]

    return {"title": title, "sections": sections}


def _parse_markdown_outline(text, original_prompt=""):
    """Parse markdown outline or structured text into {"title": str, "sections": [{"heading", "bullets", "table"}]}.

    Handles standard markdown (#, ##, ###), bold headers (**Heading**), numbered sections,
    tables, bullets, and raw paragraphs with maximum resilience.
    """
    if not text or not isinstance(text, str) or not text.strip():
        return _generate_fallback_structure(original_prompt)

    title = ""
    sections = []
    current = None

    lines = text.splitlines()

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        # Detect Document Title (# Title or Title:)
        if line.startswith('# ') and not title:
            title = line[2:].strip().strip('#*')
            continue

        # Detect Section Headings (##, ###, ####, or bold **Heading**)
        is_heading = False
        heading_text = ""

        if line.startswith(('## ', '### ', '#### ')):
            is_heading = True
            heading_text = re.sub(r'^#+\s*', '', line).strip().strip('*_')
        elif re.match(r'^(?:\d+[\.\)]\s+)?\*\*(.+?)\*\*:?$', line):
            m = re.match(r'^(?:\d+[\.\)]\s+)?\*\*(.+?)\*\*:?$', line)
            is_heading = True
            heading_text = m.group(1).strip()
        elif re.match(r'^(?:Section\s+\d+|Chapter\s+\d+|\d+\.)\s+([A-Z].+)$', line) and len(line) < 80 and not line.endswith('.'):
            is_heading = True
            heading_text = line.strip('*_')

        if is_heading and heading_text:
            if not title:
                title = heading_text
            current = {"heading": heading_text, "bullets": [], "table": []}
            sections.append(current)
            continue

        # If we haven't encountered a heading yet, start an initial Overview section
        if current is None:
            if not title and len(line) < 80 and not line.startswith(('-', '*', '|')):
                title = line.strip('*_# ')
                continue
            current = {"heading": "Overview & Summary", "bullets": [], "table": []}
            sections.append(current)

        # Detect Tables
        if line.startswith('|') and line.endswith('|'):
            if _TABLE_SEPARATOR_RE.match(line):
                continue
            cells = [c.strip() for c in line.strip('|').split('|')]
            if any(cells):
                current["table"].append(cells)
            continue

        # Detect Bullet Points
        if re.match(r'^(?:[-*+]|\d+[\.\)])\s+', line):
            clean_bullet = re.sub(r'^(?:[-*+]|\d+[\.\)])\s+', '', line).strip()
            if clean_bullet:
                current["bullets"].append(clean_bullet)
            continue

        # Regular descriptive text line
        if len(line) > 3:
            current["bullets"].append(line)

    # Normalize sections: clean empty ones and ensure either table or bullets
    valid_sections = []
    for section in sections:
        if section["table"] and len(section["table"]) >= 2:
            section["bullets"] = []
            valid_sections.append(section)
        elif section["bullets"]:
            section["table"] = None
            valid_sections.append(section)

    if not title:
        title = _clean_subject_title(original_prompt)

    if not valid_sections:
        return _generate_fallback_structure(original_prompt)

    return {"title": title or "Untitled Document", "sections": valid_sections}


def generate_document_structure(prompt, language="en"):
    """Ask the LLM for a markdown outline about `prompt` and parse it into a document structure.
    Falls back gracefully if LLM generation fails or returns empty/malformed results.
    """
    system_prompt = (
        "You write structured outlines for documents. Given a subject, respond ONLY with "
        "a markdown outline in this exact shape:\n"
        "# <Title>\n"
        "## <Section heading>\n"
        "- <bullet point>\n"
        "- <bullet point>\n"
        "## <Another section heading>\n"
        "| <column> | <column> |\n"
        "| <value> | <value> |\n\n"
        "Use bullet points for narrative sections and a markdown table only for sections "
        "that are genuinely tabular data. Include 3-6 sections. Do not include any text "
        "outside the outline."
    )
    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": prompt},
    ]

    try:
        outline_text = generate_completion(messages, language=language)
        if (
            outline_text
            and isinstance(outline_text, str)
            and "Pragna Autopilot:" not in outline_text
            and "```canvas" not in outline_text
            and len(outline_text.strip()) > 30
        ):
            parsed = _parse_markdown_outline(outline_text, original_prompt=prompt)
            if parsed and parsed.get("sections") and len(parsed["sections"]) >= 2:
                return parsed
    except Exception as exc:
        logger.warning(f"LLM completion for document outline failed, using fallback: {exc}")

    return _generate_fallback_structure(prompt, language=language)


def _build_docx(structure, filepath):
    """Build a beautifully formatted Microsoft Word document (.docx)."""
    doc = Document()
    doc.add_heading(structure.get("title") or "Document", level=0)

    for section in structure.get("sections", []):
        doc.add_heading(section.get("heading") or "Section", level=1)
        if section.get("table"):
            rows = section["table"]
            if rows and len(rows) > 0:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        if c < len(table.columns):
                            table.cell(r, c).text = str(cell)
        else:
            for bullet in section.get("bullets", []):
                doc.add_paragraph(bullet, style="List Bullet")

    doc.save(filepath)


def _build_pdf(structure, filepath):
    """Build an aesthetically styled PDF document with clean headers, tables, and typography."""
    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        leftMargin=40,
        rightMargin=40,
        topMargin=40,
        bottomMargin=40,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#0F172A'),
        alignment=0,
        spaceAfter=14,
    )

    h2_style = ParagraphStyle(
        'DocHeading2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=12,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceAfter=4,
    )

    cell_style = ParagraphStyle(
        'DocCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#1E293B'),
    )

    cell_header_style = ParagraphStyle(
        'DocCellHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#FFFFFF'),
    )

    story = [
        Paragraph(html.escape(structure.get("title") or "Document"), title_style),
        HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#E2E8F0'), spaceAfter=14),
    ]

    for section in structure.get("sections", []):
        story.append(Paragraph(html.escape(section.get("heading") or "Section"), h2_style))
        story.append(Spacer(1, 4))

        if section.get("table") and len(section["table"]) > 0:
            raw_table = section["table"]
            num_cols = max(len(r) for r in raw_table) if raw_table else 1
            available_width = 532.0  # letter width 612 - 80 margin
            col_width = available_width / max(num_cols, 1)

            formatted_table_data = []
            for r_idx, row in enumerate(raw_table):
                row_cells = []
                for c_idx in range(num_cols):
                    val = str(row[c_idx]) if c_idx < len(row) else ""
                    st = cell_header_style if r_idx == 0 else cell_style
                    row_cells.append(Paragraph(html.escape(val), st))
                formatted_table_data.append(row_cells)

            pdf_table = PdfTable(formatted_table_data, colWidths=[col_width] * num_cols)
            pdf_table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor('#1E293B')),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(pdf_table)
        else:
            for bullet in section.get("bullets", []):
                bullet_escaped = html.escape(bullet)
                story.append(Paragraph(f"&bull;&nbsp; {bullet_escaped}", body_style))

        story.append(Spacer(1, 10))

    doc.build(story)


def _build_pptx(structure, filepath):
    """Build a clean PowerPoint presentation deck (.pptx)."""
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = structure.get("title") or "Presentation"

    for section in structure.get("sections", []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = section.get("heading") or "Section"
        body = slide.placeholders[1].text_frame
        body.clear()

        lines = section.get("bullets", [])
        if not lines and section.get("table"):
            lines = [" | ".join(str(c) for c in row) for row in section["table"]]

        for i, line in enumerate(lines):
            if i == 0:
                body.text = line
            else:
                body.add_paragraph().text = line

    prs.save(filepath)


def _sanitize_filename_component(text):
    """Turn arbitrary text into a short, filesystem-safe slug."""
    cleaned = re.sub(r'[^\w\s-]', '', text or '').strip().lower()
    cleaned = re.sub(r'[\s]+', '-', cleaned)
    return cleaned[:60] or 'document'


def _sanitize_sheet_name(name):
    cleaned = re.sub(r'[\[\]:*?/\\]', '', name or '').strip()
    return cleaned[:31] or "Sheet1"


def _build_xlsx(structure, filepath):
    """Build a styled Excel workbook (.xlsx) with header highlights."""
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    table_section = next((s for s in structure.get("sections", []) if s.get("table")), None)
    if table_section and table_section["table"]:
        ws = wb.create_sheet(_sanitize_sheet_name(table_section.get("heading") or "Data"))
        for r_idx, row in enumerate(table_section["table"], start=1):
            ws.append(row)
            if r_idx == 1:
                for col_idx in range(1, len(row) + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center")
    else:
        ws = wb.create_sheet("Summary")
        ws.append(["Section", "Content"])
        for col_idx in (1, 2):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill

        for section in structure.get("sections", []):
            bullets_text = "\n".join(f"• {b}" for b in section.get("bullets", []))
            ws.append([section.get("heading", ""), bullets_text])

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

    wb.save(filepath)

