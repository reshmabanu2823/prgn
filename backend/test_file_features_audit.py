"""Comprehensive audit test for file upload analysis and document generation."""
import io
import os
import sys
import unittest
from PIL import Image

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

from app import app, MAX_UPLOAD_FILE_BYTES, _extract_text_from_blob, _extract_upload_payload


from unittest.mock import patch

SAMPLE_STRUCTURE = {
    "title": "Quarterly Financial Overview",
    "sections": [
        {"heading": "Executive Summary", "bullets": ["Revenue grew by 20%", "Operating margin improved"], "table": None},
        {"heading": "Financials", "bullets": [], "table": [["Quarter", "Revenue", "Profit"], ["Q1", "$10M", "$2M"], ["Q2", "$12M", "$2.5M"]]},
    ],
}


class FileFeaturesAuditTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        app.config['TESTING'] = True
        cls.client = app.test_client()

    def test_01_upload_plain_text(self):
        content = b"Hello from plain text upload test!"
        text, method, err = _extract_text_from_blob("sample.txt", "text/plain", content)
        self.assertIsNone(err)
        self.assertEqual(method, "plain_text")
        self.assertIn("Hello from plain text", text)

    def test_02_upload_markdown(self):
        content = b"# Title\n\n- Point A\n- Point B\n"
        text, method, err = _extract_text_from_blob("notes.md", "text/markdown", content)
        self.assertIsNone(err)
        self.assertEqual(method, "plain_text")
        self.assertIn("Point A", text)

    def test_03_upload_pdf(self):
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.drawString(100, 750, "Sample PDF Document for Upload Analysis")
        c.save()
        pdf_bytes = buf.getvalue()

        text, method, err = _extract_text_from_blob("doc.pdf", "application/pdf", pdf_bytes)
        self.assertIsNone(err)
        self.assertEqual(method, "pdf")
        self.assertIn("Sample PDF Document", text)

    def test_04_upload_docx(self):
        doc = Document()
        doc.add_heading("Docx Test Heading", level=1)
        doc.add_paragraph("This is a paragraph inside docx.")
        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        text, method, err = _extract_text_from_blob("file.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", docx_bytes)
        self.assertIsNone(err)
        self.assertEqual(method, "docx")
        self.assertIn("Docx Test Heading", text)
        self.assertIn("This is a paragraph inside docx.", text)

    def test_05_upload_xlsx(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Product", "Revenue"])
        ws.append(["Widget A", "5000"])
        buf = io.BytesIO()
        wb.save(buf)
        xlsx_bytes = buf.getvalue()

        text, method, err = _extract_text_from_blob("data.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", xlsx_bytes)
        self.assertIsNone(err)
        self.assertEqual(method, "xlsx")
        self.assertIn("Sales", text)
        self.assertIn("Widget A | 5000", text)

    def test_06_upload_pptx(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Presentation Title Slide"
        slide.placeholders[1].text = "Subtitle or bullet points in slide"
        buf = io.BytesIO()
        prs.save(buf)
        pptx_bytes = buf.getvalue()

        text, method, err = _extract_text_from_blob("slides.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation", pptx_bytes)
        self.assertIsNone(err)
        self.assertEqual(method, "pptx")
        self.assertIn("Presentation Title Slide", text)
        self.assertIn("Subtitle or bullet points in slide", text)

    def test_07_upload_image_png_jpg(self):
        img = Image.new("RGB", (100, 100), color="blue")
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        png_bytes = buf.getvalue()

        from werkzeug.datastructures import FileStorage
        fs = FileStorage(stream=io.BytesIO(png_bytes), filename="test.png", content_type="image/png")
        payload = _extract_upload_payload([fs], ["test.png"], ["image"])
        self.assertEqual(len(payload["file_summaries"]), 1)
        summary = payload["file_summaries"][0]
        self.assertTrue(summary["text_extracted"])
        self.assertEqual(summary["extraction_method"], "vision_analysis")

    def test_08_upload_oversized_file(self):
        from werkzeug.datastructures import FileStorage
        large_blob = b"A" * (MAX_UPLOAD_FILE_BYTES + 100)
        fs = FileStorage(stream=io.BytesIO(large_blob), filename="huge.txt", content_type="text/plain")
        payload = _extract_upload_payload([fs], ["huge.txt"], ["file"])
        self.assertEqual(payload["skipped_files"], 1)
        self.assertTrue(any("exceeds" in w for w in payload["warnings"]))

    @patch("services.document_generator.generate_document_structure", return_value=SAMPLE_STRUCTURE)
    def test_09_generate_and_download_docx(self, mock_gen):
        res = self.client.post("/api/documents/generate", json={
            "format": "docx",
            "prompt": "Quarterly Financial Overview",
            "language": "en"
        })
        self.assertEqual(res.status_code, 200, res.get_json())
        data = res.get_json()
        self.assertIn("download_url", data)
        self.assertIn("filename", data)
        self.assertTrue(data["filename"].endswith(".docx"))

        dl_res = self.client.get(data["download_url"])
        self.assertEqual(dl_res.status_code, 200)
        self.assertGreater(len(dl_res.data), 0)

    @patch("services.document_generator.generate_document_structure", return_value=SAMPLE_STRUCTURE)
    def test_10_generate_and_download_pdf(self, mock_gen):
        res = self.client.post("/api/documents/generate", json={
            "format": "pdf",
            "prompt": "Project Roadmap 2026",
            "language": "en"
        })
        self.assertEqual(res.status_code, 200, res.get_json())
        data = res.get_json()
        self.assertIn("download_url", data)
        self.assertTrue(data["filename"].endswith(".pdf"))

        dl_res = self.client.get(data["download_url"])
        self.assertEqual(dl_res.status_code, 200)
        self.assertGreater(len(dl_res.data), 0)

    @patch("services.document_generator.generate_document_structure", return_value=SAMPLE_STRUCTURE)
    def test_11_generate_and_download_pptx(self, mock_gen):
        res = self.client.post("/api/documents/generate", json={
            "format": "pptx",
            "prompt": "AI Strategy Pitch Deck",
            "language": "en"
        })
        self.assertEqual(res.status_code, 200, res.get_json())
        data = res.get_json()
        self.assertIn("download_url", data)
        self.assertTrue(data["filename"].endswith(".pptx"))

        dl_res = self.client.get(data["download_url"])
        self.assertEqual(dl_res.status_code, 200)
        self.assertGreater(len(dl_res.data), 0)

    @patch("services.document_generator.generate_document_structure", return_value=SAMPLE_STRUCTURE)
    def test_12_generate_and_download_xlsx(self, mock_gen):
        res = self.client.post("/api/documents/generate", json={
            "format": "xlsx",
            "prompt": "Budget Comparison Table",
            "language": "en"
        })
        self.assertEqual(res.status_code, 200, res.get_json())
        data = res.get_json()
        self.assertIn("download_url", data)
        self.assertTrue(data["filename"].endswith(".xlsx"))

        dl_res = self.client.get(data["download_url"])
        self.assertEqual(dl_res.status_code, 200)
        self.assertGreater(len(dl_res.data), 0)

    def test_13_generate_invalid_format(self):
        res = self.client.post("/api/documents/generate", json={
            "format": "exe",
            "prompt": "Malicious payload",
        })
        self.assertEqual(res.status_code, 400)


if __name__ == "__main__":
    unittest.main()
